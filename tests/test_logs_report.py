import io
from datetime import datetime

from openpyxl import load_workbook
from pypdf import PdfReader

from app.models import AuditEvent, utcnow
from tests.conftest import ADMIN, LEADER_A, login, logout


def _event(db, *, label="LOG-SEARCH-ONLY-001", status_code=200, created_at=None):
    event = AuditEvent(
        actor_user_id=1,
        actor_name="Test Admin",
        actor_role="admin",
        action="update",
        module="installation",
        entity_type="installation_record",
        entity_id="17",
        entity_label=label,
        method="POST",
        path="/installations/17/edit",
        status_code=status_code,
        ip_address="192.168.1.15",
        user_agent="pytest",
        changes_json='{"site": {"before": "Gate 1", "after": "Gate 2"}}',
        created_at=created_at or utcnow(),
    )
    db.add(event)
    db.commit()
    return event


def test_logs_report_is_admin_only_and_single_management_entry(client):
    login(client, *LEADER_A)
    assert client.get("/management/logs-report").status_code == 403

    logout(client)
    login(client, *ADMIN)
    page = client.get("/management/logs-report")
    assert page.status_code == 200
    assert 'href="/management/logs-report"' in page.text
    assert 'href="/admin/audit-log"' not in page.text
    assert 'href="/reports/technician-audit"' not in page.text
    assert "Search to view logs" in page.text
    assert client.get("/admin/audit-log").status_code == 307
    assert client.get("/reports/technician-audit").status_code == 307


def test_logs_are_hidden_until_a_filtered_search(client, db):
    event = _event(db)
    login(client, *ADMIN)

    initial = client.get("/management/logs-report")
    assert event.entity_label not in initial.text

    unfiltered = client.get("/management/logs-report?searched=1")
    assert unfiltered.status_code == 200
    assert "Choose at least one search filter" in unfiltered.text
    assert event.entity_label not in unfiltered.text

    filtered = client.get(
        "/management/logs-report",
        params={"searched": "1", "q": event.entity_label},
    )
    assert filtered.status_code == 200
    assert event.entity_label in filtered.text
    assert "192.168.1.15" in filtered.text
    assert f'/management/logs-report/events/{event.id}' in filtered.text


def test_logs_report_filters_user_system_and_failure(client, db):
    success = _event(db, label="USER-SUCCESS", status_code=200)
    failure = _event(db, label="SYSTEM-FAILURE", status_code=500)
    failure.actor_user_id = None
    failure.actor_name = "Anonymous"
    failure.actor_role = None
    db.commit()
    login(client, *ADMIN)

    user_page = client.get(
        "/management/logs-report",
        params={"searched": "1", "log_type": "user", "status": "success"},
    )
    assert success.entity_label in user_page.text
    assert failure.entity_label not in user_page.text

    system_page = client.get(
        "/management/logs-report",
        params={"searched": "1", "log_type": "system"},
    )
    assert failure.entity_label in system_page.text


def test_filtered_logs_export_pdf_and_excel(client, db):
    event = _event(db, label="EXPORT-LOG-001")
    login(client, *ADMIN)
    params = {"searched": "1", "q": event.entity_label}

    pdf = client.get("/management/logs-report/pdf", params=params)
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf.content)).pages)
    assert "Logs Report" in pdf_text
    assert event.entity_label in pdf_text

    excel = client.get("/management/logs-report/excel", params=params)
    assert excel.status_code == 200
    assert excel.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(excel.content), read_only=True)
    sheet = workbook["Logs Report"]
    assert sheet["A1"].value == "Service Management System — Logs Report"
    assert sheet["H5"].value == event.entity_label
    assert sheet["L5"].value and "Gate 2" in sheet["L5"].value


def test_logs_report_time_filter_uses_display_timezone_and_seconds(client, db):
    included = _event(db, label="AT-EXACT-SECOND", created_at=datetime(2026, 8, 25, 7, 0, 0))
    excluded = _event(db, label="ONE-SECOND-LATER", created_at=datetime(2026, 8, 25, 7, 0, 1))
    login(client, *ADMIN)
    page = client.get(
        "/management/logs-report",
        params={
            "searched": "1",
            "from_at": "2026-08-25T10:00:00",
            "to_at": "2026-08-25T10:00:00",
        },
    )
    assert page.status_code == 200
    assert included.entity_label in page.text
    assert excluded.entity_label not in page.text
    assert "10:00:00" in page.text


def test_technician_activity_is_inside_logs_report_and_search_first(client):
    login(client, *ADMIN)
    initial = client.get("/management/logs-report?view=technician")
    assert initial.status_code == 200
    assert "Select a Technical user" in initial.text
    assert "Total visits" not in initial.text

    searched = client.get(
        "/management/logs-report",
        params={"view": "technician", "searched": "1", "technician_id": "2"},
    )
    assert searched.status_code == 200
    assert "Total visits" in searched.text

    excel = client.get(
        "/management/logs-report/technician.xlsx", params={"technician_id": "2"}
    )
    assert excel.status_code == 200
    workbook = load_workbook(io.BytesIO(excel.content), read_only=True)
    assert workbook.sheetnames == ["Summary", "Work History", "Edit History"]
    assert workbook["Summary"]["A1"].value == "Technician Activity — Leader One"
